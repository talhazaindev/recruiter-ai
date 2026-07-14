import re
import fitz
import json
import os
from docx import Document
from docx.shared import Pt


HEADING_MAP = {
    # =========================
    # EDUCATION
    # =========================
    "education": "education",
    "educational background": "education",
    "academic background": "education",
    "academic qualification": "education",
    "academic qualifications": "education",
    "educational qualifications": "education",
    "qualifications": "education",
    "education & training": "education",
    "education and training": "education",

    # =========================
    # EXPERIENCE
    # =========================
    "experience": "experience",
    "work experience": "experience",
    "professional experience": "experience",
    "employment history": "experience",
    "career history": "experience",
    "work history": "experience",
    "employment": "experience",
    "professional background": "experience",
    "industry experience": "experience",
    "career experience": "experience",

    # =========================
    # SKILLS
    # =========================
    "skills": "skills",
    "technical skills": "skills",
    "core skills": "skills",
    "key skills": "skills",
    "professional skills": "skills",
    "hard skills": "skills",
    "soft skills": "skills",
    "competencies": "skills",
    "technical competencies": "skills",
    "technical expertise": "skills",
    "expertise": "skills",
    "areas of expertise": "skills",
    "technology stack": "skills",
    "tools & technologies": "skills",
    "tools and technologies": "skills",
    "technologies": "skills",
    "programming languages": "skills",
    "technical proficiencies": "skills",

    # =========================
    # PROJECTS
    # =========================
    "projects": "projects",
    "project": "projects",
    "personal projects": "projects",
    "academic projects": "projects",
    "professional projects": "projects",
    "key projects": "projects",
    "major projects": "projects",
    "selected projects": "projects",
    "relevant projects": "projects",
    "project experience": "projects",

    # =========================
    # CERTIFICATIONS
    # =========================
    "certification": "certifications",
    "certifications": "certifications",
    "certificates": "certifications",
    "licenses & certifications": "certifications",
    "licenses and certifications": "certifications",
    "professional certifications": "certifications",
    "training & certifications": "certifications",
    "training and certifications": "certifications",
    "credentials": "certifications",
    "licenses": "certifications",

    # =========================
    # PROFESSIONAL SUMMARY
    # =========================
    "summary": "summary",
    "professional summary": "summary",
    "career summary": "summary",
    "profile": "summary",
    "professional profile": "summary",
    "career profile": "summary",
    "executive summary": "summary",
    "overview": "summary",
    "about me": "summary",
    "about": "summary",
    "objective": "summary",
    "career objective": "summary",
    "professional objective": "summary",
    "personal statement": "summary",
    "career overview": "summary",

    #==============================
    #certificates
    #==============================
    # =========================
# CERTIFICATIONS
# =========================
"certification": "certifications",
"certifications": "certifications",
"certificate": "certifications",
"certificates": "certifications",
"certified courses": "certifications",
"professional certification": "certifications",
"professional certifications": "certifications",
"industry certifications": "certifications",
"technical certifications": "certifications",
"licenses": "certifications",
"license": "certifications",
"licences": "certifications",          # British spelling
"licence": "certifications",
"licenses and certifications": "certifications",
"licences and certifications": "certifications",
"credentials": "certifications",
"professional credentials": "certifications",
"training": "certifications",
"trainings": "certifications",
"training and certifications": "certifications",
"courses": "certifications",
"online courses": "certifications",
"completed courses": "certifications",
"professional development": "certifications",
"continuing education": "certifications",
"certification and training": "certifications",
}

def extract_raw_text(pdf_path):
    """
    Extract raw text from a PDF file without any formatting.
    
    Args:
        pdf_path (str): Path to the PDF file
        
    Returns:
        str: Raw text content
    """
    doc = fitz.open(pdf_path)
    all_text = ""
    
    for page in doc:
        all_text += page.get_text()
    
    doc.close()
    #print(all_text)
    return all_text

def extract_raw_text_docx(docx_path):
    """
    Extract raw text from a DOCX file, including tables.

    Args:
        docx_path (str): Path to the DOCX file

    Returns:
        str: Raw text content
    """
    doc = Document(docx_path)

    all_text = []

    # Normal paragraphs
    for para in doc.paragraphs:
        if para.text.strip():
            all_text.append(para.text)

    # Text inside tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    all_text.append(cell.text)

    return "\n".join(all_text)

def extract_spans(pdf_path):

    doc = fitz.open(pdf_path)

    spans = []

    for page_num, page in enumerate(doc, start=1):

        text_dict = page.get_text("dict",sort=True)

        for block in text_dict["blocks"]:

            if block["type"] != 0:
                continue

            for line in block["lines"]:
                #print(len(line["spans"]))
                
                for span in line["spans"]:

                    spans.append({
                        "text": span["text"],
                        "font": span["font"],
                        "size": span["size"],
                        "flags": span["flags"],
                        "bbox": span["bbox"],
                        "page": page_num,
                        "bold": span["flags"] >= 16,

                    })

    doc.close()
    #print(spans)
    return spans

def extract_docx_spans(docx_path):
    """
    Extract text runs from a DOCX file with robust font size handling.
    """
    doc = Document(docx_path)
    spans = []

    for para in doc.paragraphs:
        # Get paragraph style size as fallback
        para_size = None
        if para.style and para.style.font and para.style.font.size:
            para_size = para.style.font.size.pt

        for run in para.runs:
            text = run.text
            if not text.strip():
                continue

            # Get font name with fallback
            font_name = run.font.name
            if font_name is None:
                # Try to get from paragraph style
                if para.style and para.style.font and para.style.font.name:
                    font_name = para.style.font.name
                else:
                    font_name = "Unknown"  # Default fallback

            # Get font size with fallbacks
            size = run.font.size.pt if run.font.size else para_size or 11.0

            # Get EFFECTIVE bold status (checks style hierarchy)
            is_bold = False
            
            # 1. Check run-level explicit bold
            if run.bold is True:
                is_bold = True
            
            # 2. Check paragraph style bold
            elif para.style and para.style.font and para.style.font.bold is True:
                is_bold = True
            
            # 3. Check character style on the run
            elif hasattr(run, 'style') and run.style and run.style.font and run.style.font.bold is True:
                is_bold = True
            
            # 4. Check style hierarchy (base styles)
            else:
                style = para.style
                while style and hasattr(style, 'base_style') and style.base_style:
                    if style.base_style.font and style.base_style.font.bold is True:
                        is_bold = True
                        break
                    style = style.base_style

            spans.append({
                "text": text,
                "font": font_name,  # Always a string now
                "size": size,       # Always a float now
                "flags": None,
                "bold": is_bold,    # Now correctly detects inherited bold!
                "bbox": None,
                "page": None
            })

    #print(spans)
    return spans
# def extract_docx_spans(docx_path):
#     """
#     Extract text runs from a DOCX file in the same format as the PDF spans.

#     Args:
#         docx_path (str): Path to the DOCX file

#     Returns:
#         list: List of span dictionaries
#     """

#     doc = Document(docx_path)
#     spans = []

#     for para in doc.paragraphs:

#         for run in para.runs:

#             text = run.text
#             if not text.strip():
#                 continue

#             spans.append({
#                 "text": text,
#                 "font": run.font.name,
#                 "size": run.font.size.pt if run.font.size else None,
#                 "flags": None,
#                 "bold": bool(run.bold),
#                 "bbox": None,
#                 "page": None
#             })

#     return spans

def is_heading(span, current_font_size=None):
    text = span["text"].strip()
    font = span["font"]
    #flags = span["flags"]
    bold_flag=span["bold"]
    bold = ("Bold" in font) or (bold_flag) 


    normalized = normalize_heading(text)
    if normalized in HEADING_MAP and bold:
        return True
    

    #if not bold:
    #    return False

    if current_font_size is not None and span["size"] == current_font_size:# and normalized in HEADING_MAP:
        return True
     
    return False

def normalize_heading(text: str) -> str:
    text = text.lower().strip()
    text = text.replace("&", "and")
    text = re.sub(r"[^\w\s]", "", text)
    text = re.sub(r"\s+", " ", text)
    return text

def build_sections(spans, current_font_size=None):

    sections = {}

    current_heading = None
    current_content = []

    for span in spans:

        text = span["text"].strip()
        font = span["font"]
        #flags = span["flags"]
        if not text:
            continue

        span["text"] = text
        if is_heading(span,current_font_size):
            
            # Save previous section
            if current_heading is not None:
                if current_heading == "miscellaneous" and current_heading in sections:
                    sections[current_heading].extend(current_content)
                else:
                    sections[current_heading] = current_content
        

            normalized = normalize_heading(text)

            if normalized in HEADING_MAP:
                current_heading = HEADING_MAP[normalized]
                current_font_size = span["size"]
            else:
                #########
                current_heading="miscellaneous"
                current_font_size = span["size"]
            current_content = []
                
                
        else:
            
            if current_heading is not None:
                current_content.append({"text": text,"bold":("Bold" in span["font"]),"flags": span["flags"]})

    # Save last section
    if current_heading is not None:
        if current_heading == "miscellaneous" and current_heading in sections:
            sections[current_heading].extend(current_content)
        else:
            sections[current_heading] = current_content
        #sections[current_heading] = current_content
    #print(sections)
    return sections

def extract_document_spans(path):
    ext = os.path.splitext(path)[1].lower()

    if ext == ".pdf":
        return extract_spans(path)

    elif ext == ".docx":
        return extract_docx_spans(path)

    else:
        raise ValueError("Unsupported format")

def find_headers(path):
    #spans = extract_spans(path)
    spans=extract_document_spans(path)
    font_sizes = []

    # Find all unique font sizes
    for span in spans:
        font_sizes.append(span["size"])

    font_sizes = list(set(font_sizes))
    font_sizes.sort(reverse=True)

    sections = build_sections(spans,font_sizes[1])

    return sections
   
    """ print("Sections found:")
    for heading, content in sections.items():
        print(f"Heading: {heading}")
        for line in content:
            print(f"  {line}")
        print("-" * 40)
     """

def group_section(items):
    grouped = {}

    current_key = None

    for item in items:
        text = item["text"].strip()

        #print(item)
        if item["bold"]:#or item["flags"] >= 16:  # Assuming bold text or flagged text indicates a new key
            current_key = text
            grouped[current_key] = []

        elif current_key is not None:
            grouped[current_key].append(text)
        else:
            # No title encountered yet
            grouped.setdefault("content", []).append(text)

    return grouped

def merge_lists(data):
    if isinstance(data, dict):
        for key, value in data.items():
            data[key] = merge_lists(value)
        return data

    elif isinstance(data, list):
        # If it's a list of strings, merge into one string inside a list
        if all(isinstance(item, str) for item in data):
            merged = " ".join(item.strip() for item in data)
            return [merged]

        # Otherwise recurse
        return [merge_lists(item) for item in data]

    return data

def form_json(sections):
    with open("output.json", "w", encoding="utf-8") as f:
        json.dump(sections, f, indent=4, ensure_ascii=False)

def merge_empty_keys(data):
    """
    Merges consecutive keys with empty values into the next key
    that contains non-empty content.

    Returns a new JSON object.
    """

    new_data = {}

    for section, entries in data.items():

        # Skip sections that are not dictionaries
        if not isinstance(entries, dict):
            new_data[section] = entries
            continue

        merged = {}
        pending = []

        keys = list(entries.keys())

        for key in keys:
            value = entries[key]

            is_empty = (
                isinstance(value, list)
                and (len(value) == 0 or all(str(v).strip() == "" for v in value))
            )

            if is_empty:
                pending.append(key)
            else:
                if pending:
                    new_key = " ".join(pending + [key])
                    pending.clear()
                else:
                    new_key = key

                merged[new_key] = value

        # Handle trailing empty keys
        if pending:
            if merged:
                last_key = next(reversed(merged))
                new_last_key = last_key + " " + " ".join(pending)
                merged[new_last_key] = merged.pop(last_key)
            else:
                merged[" ".join(pending)] = [""]

        new_data[section] = merged

    return new_data

def parse_cv(path):
    print("Parsing CV from path:", path)
    raw_text = extract_raw_text(path)
    sections = find_headers(path)
    #print(sections)
    for section_name, items in sections.items():
        sections[section_name] = group_section(items)
    sections["raw_text"] = raw_text
    merged_sections = merge_lists(sections)
    merged_sections = merge_empty_keys(merged_sections)
    form_json(merged_sections)
    #return raw_text
    